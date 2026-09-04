"""Approval-gated, in-memory Word document for a reviewed draft package."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
import re
from typing import Any, Literal

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, ConfigDict, Field

from nerc_compliance_intelligence.case_intake import CaseIntakeAssessment


WORD_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_INVALID_XML_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
_NAVY = "17365D"
_BLUE = "2E74B5"
_PALE_BLUE = "EAF2F8"
_PALE_GRAY = "F2F4F7"
_DARK_GRAY = "404040"


class ExportDecision(BaseModel):
    """The explicit human decision that authorizes preparing one Word package."""

    model_config = ConfigDict(frozen=True)

    decision: Literal["approve"] = "approve"
    reviewer_role: str = Field(min_length=1, max_length=120)
    rationale: str = Field(min_length=1, max_length=500)
    recorded_at: datetime


class ExportSource(BaseModel):
    """One public source excerpt and its traceability metadata."""

    requirement_reference: str
    locator: str
    excerpt: str
    source_url: str | None = None
    retrieval_date: str | None = None
    verified_local_match: bool


class ExportRemediationStep(BaseModel):
    step_number: int = Field(ge=1)
    action: str
    owner_role: str
    decision: str
    end_state: str


class ExportRequirementPackage(BaseModel):
    """Source, control, and remediation content for one requirement."""

    requirement_reference: str
    sources: list[ExportSource]
    control_id: str
    control_title: str
    objective: str
    activities: list[str]
    control_type: str
    owner_role: str
    performer: str
    frequency: str
    procedure: str
    evidence_expectation: str
    exception_escalation: str
    test_procedure: str
    assumptions: list[str]
    tailoring_questions: list[str]
    remediation_steps: list[ExportRemediationStep]


class ReviewPackageExport(BaseModel):
    """Complete typed content used to build the Word package."""

    app_name: str = "CIP Wayfinder"
    document_title: str = "Draft control and remediation review package"
    source_file_name: str
    standard_version: str
    functional_entities: list[str]
    regional_entities: list[str]
    asset_scope: list[str]
    review_objective: str
    requirements: list[ExportRequirementPackage] = Field(min_length=1)
    export_decision: ExportDecision


def _clean_text(value: object) -> str:
    """Remove characters that are invalid in OOXML while preserving readable text."""
    return _INVALID_XML_CHARACTERS.sub("", str(value)).strip()


def _export_remediation_steps(remediation: Any) -> list[ExportRemediationStep]:
    return [
        ExportRemediationStep(
            step_number=step.step_number,
            action=_clean_text(step.action.text),
            owner_role=_clean_text(step.owner_role.text),
            decision=_clean_text(step.decision.text),
            end_state=_clean_text(step.end_state.text),
        )
        for step in remediation.steps
    ]


def build_review_package_export(
    data: dict[str, Any],
    assessment: CaseIntakeAssessment | None,
    export_decision: ExportDecision,
) -> ReviewPackageExport:
    """Convert the reviewed Streamlit package into a strict export contract."""
    uploaded = data["uploaded"]
    requirement_packages: list[ExportRequirementPackage] = []
    for package in data["packages"]:
        mapping = package["mapping"]
        control = package["control"]
        remediation = package["remediation"]
        sources = [
            ExportSource(
                requirement_reference=mapping.requirement_reference,
                locator=f"page {chunk.metadata.page}, {chunk.metadata.section}",
                excerpt=_clean_text(chunk.text),
                source_url=_clean_text(chunk.metadata.source_url),
                retrieval_date=chunk.metadata.retrieval_date.isoformat(),
                verified_local_match=True,
            )
            for chunk in package["source_chunks"]
        ]
        if not sources:
            sources.append(
                ExportSource(
                    requirement_reference=mapping.requirement_reference,
                    locator=_clean_text(mapping.source_locator),
                    excerpt="No matching approved local source excerpt was available. Verify the requirement text before relying on this draft.",
                    verified_local_match=False,
                )
            )
        requirement_packages.append(
            ExportRequirementPackage(
                requirement_reference=mapping.requirement_reference,
                sources=sources,
                control_id=control.control_id,
                control_title=_clean_text(control.title.text),
                objective=_clean_text(control.objective.text),
                activities=[_clean_text(item.activity.text) for item in control.activities],
                control_type=_clean_text(control.control_type.text),
                owner_role=_clean_text(control.owner_role.text),
                performer=_clean_text(control.performer.text),
                frequency=_clean_text(control.trigger_frequency.text),
                procedure=_clean_text(control.procedure.text),
                evidence_expectation=_clean_text(control.evidence_expectation.text),
                exception_escalation=_clean_text(control.exception_escalation.text),
                test_procedure=_clean_text(control.test_procedure.text),
                assumptions=[_clean_text(item.text) for item in control.assumptions],
                tailoring_questions=[_clean_text(item.text) for item in control.tailoring_questions],
                remediation_steps=_export_remediation_steps(remediation),
            )
        )

    intake = assessment.intake if isinstance(assessment, CaseIntakeAssessment) else None
    return ReviewPackageExport(
        source_file_name=_clean_text(uploaded.file_name),
        standard_version=f"{uploaded.standard.standard_id}-{uploaded.standard.version}",
        functional_entities=list(intake.functional_entity) if intake else ["Not recorded"],
        regional_entities=list(intake.jurisdiction) if intake else ["Not recorded"],
        asset_scope=list(intake.asset_scope) if intake else ["Not recorded"],
        review_objective=_clean_text(intake.review_objective) if intake and intake.review_objective else "Not recorded",
        requirements=requirement_packages,
        export_decision=export_decision,
    )


def _set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(cell: Any, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _set_repeat_table_header(row: Any) -> None:
    row_properties = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    row_properties.append(table_header)


def _set_table_geometry(table: Any, widths_dxa: tuple[int, ...], indent_dxa: int = 120) -> None:
    """Set exact Word table, grid, and cell widths for predictable rendering."""
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table_properties.append(table_width)
    table_width.set(qn("w:w"), str(sum(widths_dxa)))
    table_width.set(qn("w:type"), "dxa")

    table_indent = table_properties.first_child_found_in("w:tblInd")
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        table_properties.append(table_indent)
    table_indent.set(qn("w:w"), str(indent_dxa))
    table_indent.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for grid_column in list(grid):
        grid.remove(grid_column)
    for width in widths_dxa:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa, strict=True):
            cell_width = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(cell_width)
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")


def _set_run_font(run: Any, name: str = "Calibri", size: float | None = None, color: str | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def _add_labeled_paragraph(document: WordDocument, label: str, value: str, *, after: float = 5) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(after)
    label_run = paragraph.add_run(f"{label}: ")
    _set_run_font(label_run, bold=True, color=_DARK_GRAY)
    value_run = paragraph.add_run(_clean_text(value))
    _set_run_font(value_run)


def _add_callout(document: WordDocument, title: str, text: str, fill: str = _PALE_BLUE) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_shading(cell, fill)
    _set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    title_run = paragraph.add_run(f"{title}\n")
    _set_run_font(title_run, bold=True, color=_NAVY)
    text_run = paragraph.add_run(_clean_text(text))
    _set_run_font(text_run, size=10.5)
    _set_table_geometry(table, (9360,), indent_dxa=180)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _configure_document(document: WordDocument) -> None:
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size, color, before, after in (
        ("Title", 23, _NAVY, 0, 6),
        ("Heading 1", 16, _BLUE, 16, 8),
        ("Heading 2", 13, _BLUE, 12, 6),
        ("Heading 3", 11.5, _NAVY, 8, 4),
    ):
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "CIP Wayfinder | Draft review package"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_run_font(header.runs[0], size=8.5, color="6B7280")
    footer = section.footer.paragraphs[0]
    footer.text = "Draft guidance | Human review required | Not a compliance determination"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(footer.runs[0], size=8, color="6B7280")


def _add_metadata_table(document: WordDocument, package: ReviewPackageExport) -> None:
    rows = [
        ("Standard", package.standard_version),
        ("Source document", package.source_file_name),
        ("Functional entities", ", ".join(package.functional_entities)),
        ("Regional entities", ", ".join(package.regional_entities)),
        ("Asset scope", ", ".join(package.asset_scope)),
        ("Review objective", package.review_objective),
        ("Package decision", f"Approved for download or email by {package.export_decision.reviewer_role}"),
        ("Decision recorded", package.export_decision.recorded_at.isoformat()),
    ]
    table = document.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row, (label, value) in zip(table.rows, rows, strict=True):
        row.cells[0].width = Inches(1.65)
        row.cells[1].width = Inches(4.85)
        for cell in row.cells:
            _set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_shading(row.cells[0], _PALE_GRAY)
        label_run = row.cells[0].paragraphs[0].add_run(label)
        _set_run_font(label_run, size=10, bold=True, color=_DARK_GRAY)
        value_run = row.cells[1].paragraphs[0].add_run(_clean_text(value))
        _set_run_font(value_run, size=10)
    _set_table_geometry(table, (2376, 6984))


def _add_requirement_summary(document: WordDocument, package: ReviewPackageExport) -> None:
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (0.8, 1.15, 2.7, 1.85)
    headers = ("Req.", "Source", "Draft control", "Owner")
    header = table.rows[0]
    _set_repeat_table_header(header)
    for cell, width, label in zip(header.cells, widths, headers, strict=True):
        cell.width = Inches(width)
        _set_cell_shading(cell, _NAVY)
        _set_cell_margins(cell)
        run = cell.paragraphs[0].add_run(label)
        _set_run_font(run, size=9.5, bold=True, color="FFFFFF")
    for item in package.requirements:
        row = table.add_row()
        values = (
            item.requirement_reference,
            "Matched" if any(source.verified_local_match for source in item.sources) else "Verify",
            item.control_title,
            item.owner_role,
        )
        for cell, width, value in zip(row.cells, widths, values, strict=True):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            run = cell.paragraphs[0].add_run(_clean_text(value))
            _set_run_font(run, size=9.5)
    _set_table_geometry(table, (1152, 1656, 3888, 2664))


def _add_requirement_detail(document: WordDocument, item: ExportRequirementPackage) -> None:
    document.add_heading(f"{item.requirement_reference} - {item.control_title}", level=2)
    document.add_heading("Sources", level=3)
    for source_number, source in enumerate(item.sources, start=1):
        status = "Approved local corpus match" if source.verified_local_match else "Source verification required"
        _add_labeled_paragraph(document, f"Source {source_number}", f"{status} | {source.locator}", after=2)
        if source.retrieval_date:
            _add_labeled_paragraph(document, "Retrieved", source.retrieval_date, after=2)
        if source.source_url:
            _add_labeled_paragraph(document, "Source URL", source.source_url, after=4)
        _add_callout(document, "Source excerpt", source.excerpt, fill=_PALE_GRAY)

    document.add_heading("Draft control", level=3)
    _add_labeled_paragraph(document, "Control ID", item.control_id)
    _add_labeled_paragraph(document, "Objective", item.objective)
    _add_labeled_paragraph(document, "Control type", item.control_type)
    _add_labeled_paragraph(document, "Owner role", item.owner_role)
    _add_labeled_paragraph(document, "Performer", item.performer)
    _add_labeled_paragraph(document, "Frequency or trigger", item.frequency)
    _add_labeled_paragraph(document, "Procedure", item.procedure)
    _add_labeled_paragraph(document, "Evidence expectation", item.evidence_expectation)
    _add_labeled_paragraph(document, "Exception and escalation", item.exception_escalation)
    _add_labeled_paragraph(document, "Test procedure", item.test_procedure)
    document.add_paragraph("Activities", style="Heading 3")
    for activity in item.activities:
        document.add_paragraph(activity, style="List Bullet")
    document.add_paragraph("Assumptions to verify", style="Heading 3")
    for assumption in item.assumptions:
        document.add_paragraph(assumption, style="List Bullet")
    document.add_paragraph("Tailoring questions", style="Heading 3")
    for question in item.tailoring_questions:
        document.add_paragraph(question, style="List Bullet")

    document.add_heading("Draft remediation", level=3)
    for step in item.remediation_steps:
        step_heading = document.add_paragraph(style="Heading 3")
        step_heading.add_run(f"Step {step.step_number}")
        _add_labeled_paragraph(document, "Action", step.action)
        _add_labeled_paragraph(document, "Owner role", step.owner_role)
        _add_labeled_paragraph(document, "Decision point", step.decision)
        _add_labeled_paragraph(document, "Expected end state", step.end_state, after=10)


def render_review_package_docx(package: ReviewPackageExport) -> bytes:
    """Render a polished Word package entirely in memory after human approval."""
    document = Document()
    _configure_document(document)
    document.core_properties.title = f"{package.standard_version} draft review package"
    document.core_properties.subject = "Source-grounded draft controls and remediation for human review"
    document.core_properties.author = package.app_name

    title = document.add_paragraph(style="Title")
    title.add_run(package.app_name)
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle_run = subtitle.add_run(package.document_title)
    _set_run_font(subtitle_run, size=14, color=_DARK_GRAY)
    _add_callout(
        document,
        "What approval means",
        "Approved as a draft package for download or email delivery. This is not a compliance determination, legal opinion, operational authorization, or approval to implement a control without organization-specific SME review.",
    )
    _add_metadata_table(document, package)
    document.add_paragraph()
    _add_labeled_paragraph(document, "Reviewer rationale", package.export_decision.rationale, after=10)

    document.add_heading("Package summary", level=1)
    _add_requirement_summary(document, package)
    document.add_heading("Requirements, sources, controls, and remediation", level=1)
    for requirement in package.requirements:
        _add_requirement_detail(document, requirement)

    document.add_heading("Human package decision", level=1)
    _add_labeled_paragraph(document, "Decision", "Approved for download or email delivery")
    _add_labeled_paragraph(document, "Reviewer role", package.export_decision.reviewer_role)
    _add_labeled_paragraph(document, "Rationale", package.export_decision.rationale)
    _add_labeled_paragraph(document, "Recorded at", package.export_decision.recorded_at.isoformat())
    _add_callout(
        document,
        "Required next step",
        "An authorized SME must tailor the draft to the organization, confirm applicability and effective versions, validate each citation, and use a separately governed process before implementation.",
        fill="FFF4CE",
    )

    output = BytesIO()
    document.save(output)
    return output.getvalue()
