"""Tests for the approval-gated, in-memory Word package."""

from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZipFile

from docx import Document

from nerc_compliance_intelligence.app import demo_dashboard_data
from nerc_compliance_intelligence.package_export import ExportDecision, build_review_package_export, render_review_package_docx


def _decision() -> ExportDecision:
    return ExportDecision(
        reviewer_role="CIP compliance manager",
        rationale="Approved as a draft package for controlled local review only.",
        recorded_at=datetime(2026, 9, 2, 12, 0, tzinfo=timezone.utc),
    )


def _document_text(docx_bytes: bytes) -> str:
    document = Document(BytesIO(docx_bytes))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def test_build_review_package_preserves_requirements_and_draft_boundary() -> None:
    data = demo_dashboard_data()

    package = build_review_package_export(data, None, _decision())

    assert package.standard_version == "CIP-010-5"
    assert package.source_document_type == "NERC Reliability Standard"
    assert package.referenced_standards == ["CIP-010-5"]
    assert package.review_objective == "Prepare a source-grounded draft package for SME tailoring."
    assert package.requirements[0].requirement_reference == "Document overview"
    assert package.requirements[0].domain == "Requirement"
    assert package.requirements[0].applicable_systems
    assert package.requirements[0].plain_language_summary
    assert package.requirements[0].sources[0].verified_local_match is False
    assert package.requirements[0].sources[0].source_status == "Content-validated uploaded NERC document"
    assert package.export_decision.decision == "approve"


def test_render_review_package_creates_a_valid_word_file() -> None:
    package = build_review_package_export(demo_dashboard_data(), None, _decision())

    docx_bytes = render_review_package_docx(package)

    assert docx_bytes.startswith(b"PK")
    with ZipFile(BytesIO(docx_bytes)) as archive:
        assert archive.testzip() is None
        assert "word/document.xml" in archive.namelist()
    text = _document_text(docx_bytes)
    assert "CIP Wayfinder" in text
    assert "Knowledge items, sources, controls, and remediation" in text
    assert "Synthetic CIP-010-5 learning document" in text
    assert "Approved as a draft package for download or email delivery" in text
    assert "not a compliance determination" in text.lower()
    assert "CIP compliance manager" in text
    assert "Human package decision" in text
    assert "Review Package Agent objective" in text
