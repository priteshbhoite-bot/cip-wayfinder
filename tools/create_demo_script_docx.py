"""Create the Word version of the CIP Wayfinder recording script.

This small builder keeps the deliverable reproducible.  It does not read
credentials, evidence, or the local corpus; all content is safe demo guidance.
"""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("docs/CIP-Wayfinder-Demo-Script.docx")
NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F4F6F9"


def set_run_font(run, size: float, color: str = NAVY, bold: bool = False, italic: bool = False) -> None:
    """Apply a predictable Calibri font to a run."""
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), value)
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[float]) -> None:
    """Apply the compact-reference-guide fixed table layout."""
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.LEFT
    tbl_pr = table._tbl.tblPr
    table_width = OxmlElement("w:tblW")
    table_width.set(qn("w:w"), "9360")
    table_width.set(qn("w:type"), "dxa")
    tbl_pr.append(table_width)
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    indent = OxmlElement("w:tblInd")
    # A zero-indent override keeps full-width timeline tables inside the page
    # margins in Word and LibreOffice while preserving the 6.5-inch grid.
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    tbl_pr.append(indent)
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
        grid_col = table._tbl.tblGrid.gridCol_lst[index]
        grid_col.set(qn("w:w"), str(round(width * 1440)))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)


def style_paragraph(paragraph, after: float = 6, before: float = 0, line: float = 1.25) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    style_paragraph(paragraph, after=4)
    run = paragraph.add_run(text)
    set_run_font(run, 11)


def add_heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    set_run_font(run, {1: 16, 2: 13, 3: 12}[level], {1: BLUE, 2: BLUE, 3: DARK_BLUE}[level], bold=True)


def add_body(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph)
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_run_font(run, 11, bold=True)
        run = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(run, 11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, 11)


def add_callout(doc: Document, title: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    shade_cell(cell, PALE_GRAY)
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, after=2)
    run = paragraph.add_run(title)
    set_run_font(run, 11, DARK_BLUE, bold=True)
    paragraph = cell.add_paragraph()
    style_paragraph(paragraph, after=0)
    run = paragraph.add_run(text)
    set_run_font(run, 10.5, NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_timeline_table(doc: Document, rows: list[tuple[str, str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    set_table_geometry(table, [0.72, 2.32, 3.46])
    headers = ("Time", "Show", "Say")
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        style_paragraph(p, after=0, line=1.0)
        run = p.add_run(header)
        set_run_font(run, 10, DARK_BLUE, bold=True)
    for time, show, say in rows:
        cells = table.add_row().cells
        for index, value in enumerate((time, show, say)):
            p = cells[index].paragraphs[0]
            style_paragraph(p, after=0, line=1.1)
            run = p.add_run(value)
            set_run_font(run, 9.2, NAVY, bold=(index == 0))


def add_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(paragraph, after=0, line=1.0)
    run = paragraph.add_run("CIP Wayfinder | Learning MVP | Page ")
    set_run_font(run, 8.5, MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def build_document() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)

    # Workshop-agenda header pattern: a concise title stack and timing strip.
    kicker = doc.add_paragraph()
    style_paragraph(kicker, after=1)
    run = kicker.add_run("RECORDED DEMONSTRATION")
    set_run_font(run, 10, BLUE, bold=True)
    title = doc.add_paragraph()
    style_paragraph(title, after=5)
    run = title.add_run("CIP Wayfinder")
    set_run_font(run, 29, NAVY, bold=True)
    subtitle = doc.add_paragraph()
    style_paragraph(subtitle, after=14)
    run = subtitle.add_run("Under-five-minute demo script and recording checklist")
    set_run_font(run, 13, MUTED)

    strip = doc.add_table(rows=1, cols=3)
    set_table_geometry(strip, [2.16, 2.16, 2.18])
    for cell, label, value in zip(
        strip.rows[0].cells,
        ("TARGET LENGTH", "DEMO USER", "SAFETY BOUNDARY"),
        ("4 minutes 40 seconds", "Compliance analyst", "Local + synthetic"),
    ):
        shade_cell(cell, PALE_BLUE)
        p = cell.paragraphs[0]
        style_paragraph(p, after=1, line=1.0)
        r = p.add_run(label + "\n")
        set_run_font(r, 8.5, DARK_BLUE, bold=True)
        r = p.add_run(value)
        set_run_font(r, 10, NAVY, bold=True)

    add_heading(doc, "Purpose", 1)
    add_body(doc, "Use this script to record a concise product demonstration for the Week 3 submission. It proves the end-to-end learning MVP: local source review, draft control design, synthetic evidence and baseline analysis, bounded graph control flow, human approval, guarded export, tests, and evaluations.")
    add_callout(doc, "Important wording", "Call the results drafts and review observations. Do not say the app makes a NERC compliance determination, gives legal advice, or changes an operational system.")

    add_heading(doc, "Before you record", 1)
    for item in (
        "Start the Streamlit app and open the local address in your browser.",
        "Use only the approved public CIP-010-5 or CIP-007-6 source and the fictional Northstar Grid Services case.",
        "Keep .env files, API keys, terminal history, and non-synthetic documents out of the recording.",
        "Prepare the local evaluation command: uv run python -m nerc_compliance_intelligence.evaluations.",
        "If showing the optional Kimi-K3 draft, show an already-approved result or the approval boundary—never a secret or a paid request without approval.",
    ):
        add_bullet(doc, item)

    doc.add_page_break()
    add_heading(doc, "Timed recording script", 1)
    add_body(doc, "Read naturally rather than word-for-word. The “Show” column tells you what to click; the “Say” column gives the key message for the reviewer.")
    rows = [
        ("0:00–0:25", "Landing page", "CIP Wayfinder helps a compliance analyst explore a public NERC CIP standard and build a starting-point review package. Northstar Grid Services, SUB-ALPHA-RTU-01, and all evidence are fictional or synthetic. The app does not declare compliance and does not make operational changes."),
        ("0:25–0:55", "Start a review", "Enter the Functional Entity, jurisdiction, asset scope, review objective, and upload one approved public standard. First leave a required field empty. The Applicability Agent asks for the missing detail instead of guessing. Complete it and continue."),
        ("0:55–1:20", "Review package → Source and requirement", "The package now shows all requirements found for this uploaded standard. Each material claim is tied to a local source with its standard/version, page or section, source URL, and retrieval date. The local corpus is read-only."),
        ("1:20–1:50", "Graph progress", "Here is a safe failure and recovery. A missing requirement reference follows the empty-retrieval route, performs no more than one repair, and stops safely when it cannot prove a source. I then rerun with a known requirement reference; the system does not invent one."),
        ("1:50–2:25", "Controls, Evidence, Baseline", "The app produces draft controls linked to retrieved requirement IDs. Synthetic evidence is treated as untrusted input; the app extracts observable facts, gaps, age, and SME questions. The baseline comparison turns only bounded differences into review observations."),
        ("2:25–2:50", "Findings and remediation", "Risk is a simple High, Medium, or Low demo rubric based on evidence completeness, asset variance, and requirement criticality. Potential findings and remediation steps remain drafts. Nothing is automatically closed."),
        ("2:50–3:20", "Review graph progress", "The visual path explains the tested LangGraph flow: missing scope, empty retrieval, tool error retries limited to two, validation repair limited to one, and a safe stop. The path reaches a human approval interrupt before any export."),
        ("3:20–3:45", "Human decision", "I can choose edit, reject, or approve. I will edit once, give a reviewer role and reason, and return to the same approval interrupt using the saved thread ID and local checkpointer."),
        ("3:45–4:15", "Approve and guarded export", "With explicit reviewer data, the approved graph route may create a synthetic JSON workflow package only under the safe local outputs folder. It does not create a ticket, contact an external system, or change an operational control."),
        ("4:15–4:40", "Terminal → evaluations", "Finally, the offline evaluation runner reports 15 passing checks. The limits remain clear: human tailoring is required, optional tracing is disconnected by default, and this MVP is informational—not a compliance conclusion."),
    ]
    add_timeline_table(doc, rows)

    doc.add_page_break()
    add_heading(doc, "Optional source-grounded model moment", 1)
    add_body(doc, "Use this only if you already have an approved public-source result. It is not needed to prove the offline MVP.")
    add_callout(doc, "Show", "Open Review and send a source-grounded draft. Point to the selected requirement, cited local excerpts, model name, bounded token limit, timeout, and the human approval step before a paid external call.")
    add_callout(doc, "Say", "The live adapter is isolated behind a provider boundary. It receives the selected public requirement and local citations, returns a structured draft, and cannot export or change an operational system. The fake provider remains the automated-test default.")

    add_heading(doc, "Close-out checklist", 1)
    for item in (
        "Keep the final video under five minutes.",
        "Show failure, recovery, interrupt, one human edit or approval, and the guarded local export boundary.",
        "Show citations and provenance for material NERC claims.",
        "End with tests/evaluations and the no-compliance-declaration limitation.",
        "Do not show secrets, confidential records, or live utility asset exports.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "Recording fallback", 2)
    add_body(doc, "If a live model call is slow or unavailable, skip it. The deterministic fake-provider and graph demonstration still prove tools, state, control flow, errors, retries, checkpointer, interrupt, human review, tracing boundary, evaluation, and end-to-end package behavior.")

    properties = doc.core_properties
    properties.title = "CIP Wayfinder Demo Script"
    properties.subject = "Under-five-minute product demonstration"
    properties.author = "CIP Wayfinder project team"
    properties.comments = "Local learning MVP; no confidential data or credentials."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(f"Created {OUTPUT}")
